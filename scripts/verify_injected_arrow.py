"""Real intervention experiment. Requires a completed task with a native directed connector.

The original run remains immutable. The altered candidate is freshly rendered and judged;
only the actual critic feedback is supplied to a new Astra repair invocation.
"""
import argparse,asyncio,shutil,sys,json
from pathlib import Path
from pptbench_harness.config import load_config,ROOT,environment
from pptbench_harness.workspace import atomic_json,read_json,sha
from pptbench_harness.checkpoints import Budget
from pptbench_harness.codex_runner import invoke
from pptbench_harness.models import AuthorResult
from pptbench_harness.executor import execute_build,validate_and_render
from pptbench_harness.critic import judge_candidate
from pptbench_harness.feedback import diagnose
from pptbench_harness.orchestrator import source_files
from pptx import Presentation
from pptx.oxml.ns import qn

def select_arrow(pptx):
    def walk(shapes):
        for s in shapes:
            if s.shape_type==6:yield from walk(s.shapes)
            else:yield s
    for s in walk(Presentation(pptx).slides[0].shapes):
        ln=s._element.find('.//'+qn('a:ln'))
        if ln is None:continue
        h=ln.find(qn('a:headEnd'));t=ln.find(qn('a:tailEnd'))
        ht=h.get('type','none') if h is not None else 'none';tt=t.get('type','none') if t is not None else 'none'
        if (ht!='none') != (tt!='none'):return s.shape_id,s.name,ht,tt
    raise RuntimeError('No native single-direction arrow found; choose another real task, do not inject a fake benchmark input.')

async def main(a):
    c=load_config();original=Path(a.task_dir).resolve();result=read_json(original/'final/result.json')
    parent=original/'candidates'/result['selected_candidate']
    sid,name,head,tail=select_arrow(parent/'validated/reconstruction.pptx')
    experiment=Path(c['output']['root'])/('injected-'+original.name+'-'+__import__('uuid').uuid4().hex[:6])
    experiment.mkdir();input_dir=experiment/'input';shutil.copytree(original/'input',input_dir)
    budget=Budget(experiment/'budget.json',c['loop']);candidate=experiment/'candidates/000';source=candidate/'source';source.mkdir(parents=True)
    for p in source_files(parent/'source'):
        dest=source/p.relative_to(parent/'source');dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,dest)
    (source/'build.py').rename(source/'build_original.py')
    wrapper='''import runpy,sys\nfrom pathlib import Path\nfrom pptx import Presentation\nfrom pptx.oxml.xmlchemy import OxmlElement\nfrom pptx.oxml.ns import qn\nrunpy.run_path(str(Path(__file__).with_name('build_original.py')),run_name='__main__')\nout=Path(sys.argv[sys.argv.index('--output')+1])\np=Presentation(out)\ndef walk(shapes):\n for s in shapes:\n  if s.shape_type==6: yield from walk(s.shapes)\n  else: yield s\nfor s in walk(p.slides[0].shapes):\n if s.shape_id==TARGET_ID:\n  ln=s._element.find('.//'+qn('a:ln'))\n  for tag,value in [('headEnd',TAIL),('tailEnd',HEAD)]:\n   e=ln.find(qn('a:'+tag))\n   if e is None:e=OxmlElement('a:'+tag);ln.append(e)\n   e.set('type',value)\np.save(out)\n'''
    wrapper=wrapper.replace('TARGET_ID',str(sid)).replace('TAIL',repr(tail)).replace('HEAD',repr(head))
    (source/'build.py').write_text(wrapper)
    atomic_json(experiment/'intervention.json',{'injected_failure':True,'task_id':original.name,'original_task_dir':str(original),'shape_id':sid,'object_id':name,'before':{'head':head,'tail':tail},'after':{'head':tail,'tail':head},'official_comparable':False})
    await execute_build(source,candidate/'raw',c,budget)
    artifact=await validate_and_render(candidate,input_dir,c,budget)
    if not artifact['valid']:raise RuntimeError('Injection produced invalid artifact')
    before=await judge_candidate(candidate,original.name,input_dir,c,budget)
    feedback=await diagnose(candidate,input_dir,c,budget)
    repaired=experiment/'candidates/001';dst=repaired/'source';dst.mkdir(parents=True)
    for p in source_files(source):
        target=dst/p.relative_to(source);target.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,target)
    atomic_json(dst/'feedback.json',feedback)
    await invoke(c,budget,dst,(ROOT/'prompts/repair.txt').read_text(),AuthorResult.model_json_schema(),[dst/'input/reference.png'],AuthorResult.model_validate)
    build=await execute_build(dst,repaired/'raw',c,budget)
    if not build['valid_output']:raise RuntimeError('Repair build failed')
    artifact=await validate_and_render(repaired,input_dir,c,budget)
    after=await judge_candidate(repaired,original.name,input_dir,c,budget) if artifact['valid'] else None
    # Verification is independent of the author claim: inspect actual OOXML endpoints.
    repaired_shape=None
    def walk(shapes):
        for s in shapes:
            if s.shape_type==6:yield from walk(s.shapes)
            else:yield s
    for s in walk(Presentation(repaired/'validated/reconstruction.pptx').slides[0].shapes):
        if s.name==name:
            ln=s._element.find('.//'+qn('a:ln'))
            if ln is not None:
                repaired_shape={}
                for key,tag in [('head','headEnd'),('tail','tailEnd')]:
                    e=ln.find(qn('a:'+tag));repaired_shape[key]=e.get('type','none') if e is not None else 'none'
    atomic_json(experiment/'result.json',{'injected_failure':True,'before':before,'after':after,'repaired_arrow':repaired_shape,
      'arrow_restored':repaired_shape=={'head':head,'tail':tail},'agent_calls':budget.calls,'elapsed_seconds':budget.elapsed,'cost_usd':None})
    print(experiment)

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--task-dir',required=True);asyncio.run(main(p.parse_args()))
