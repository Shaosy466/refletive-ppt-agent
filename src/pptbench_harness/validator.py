from pathlib import Path
import sys
from eval_harness.pptx import validate_reconstruction_pptx,render_pptx,PptxError
from .workspace import read_json,atomic_json

def verify_object_map(pptx_path, map_path):
    from pptx import Presentation
    def walk(shapes):
        for shape in shapes:
            yield shape
            if shape.shape_type == 6: yield from walk(shape.shapes)
    actual={shape.shape_id:shape for shape in walk(Presentation(pptx_path).slides[0].shapes)}
    rows=read_json(map_path)
    if not isinstance(rows,list) or not rows: raise PptxError('object-map must be a nonempty list')
    seen=set();canonical=set()
    for row in rows:
        required={'object_id','pptx_shape_id','kind','bbox','text','source_file','source_symbol'}
        if not isinstance(row,dict) or not required.issubset(row): raise PptxError('Incomplete object-map row')
        shape=actual.get(row['pptx_shape_id'])
        if shape is None: raise PptxError('object-map references nonexistent shape')
        if shape.name==row['object_id'] or shape.name.startswith(row['object_id']+':'):
            canonical.add(shape.shape_id)
        seen.add(shape.shape_id)
    if set(actual)-seen: raise PptxError('object-map omits native shapes')
    if set(actual)-canonical: raise PptxError('object-map lacks canonical shape names')
    return {'mapped_shapes':len(seen),'actual_shapes':len(actual)}

def main():
    j=read_json(Path(sys.argv[1]));target=Path(j['report'])
    try:
        report=validate_reconstruction_pptx(Path(j['pptx']),reference_path=Path(j['reference']),
            approved_resource_paths=[Path(p) for p in j['assets']],enforce_resource_integrity=True,enforce_image_similarity=True)
        if j.get('require_map'):
            report['object_map_audit']=verify_object_map(Path(j['pptx']),Path(j['object_map']))
    except (PptxError,ValueError,KeyError) as exc:
        atomic_json(target,{'valid':False,'stage':'artifact','reason':str(exc)});return 3
    try:render_pptx(Path(j['pptx']),Path(j['screenshot']),soffice_command=j['soffice'])
    except Exception as exc:
        atomic_json(target,{'valid':False,'infra_error':f'Render failed: {exc}'});return 4
    atomic_json(target,{'valid':True,'components':report});return 0
if __name__=='__main__':raise SystemExit(main())
