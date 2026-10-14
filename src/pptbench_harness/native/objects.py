"""Small native helpers; all coordinates normalized to the slide canvas."""
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
import json

class NativeSlide:
    def __init__(self,aspect=16/9):
        self.prs=Presentation();self.prs.slide_width=Inches(13.333)
        self.prs.slide_height=round(self.prs.slide_width/aspect)
        self.slide=self.prs.slides.add_slide(self.prs.slide_layouts[6]);self.objects=[]
    def box(self,b): return [round(v*(self.prs.slide_width if i%2==0 else self.prs.slide_height)) for i,v in enumerate(b)]
    def register(self,shape,id,kind,symbol='build'):
        shape.name=id
        self.objects.append({'object_id':id,'pptx_shape_id':shape.shape_id,'kind':kind,
            'bbox':[shape.left/self.prs.slide_width,shape.top/self.prs.slide_height,shape.width/self.prs.slide_width,shape.height/self.prs.slide_height],
            'text':shape.text if shape.has_text_frame else '', 'source_file':'build.py','source_symbol':symbol})
        return shape
    def rect(self,id,bbox,text='',fill='FFFFFF',line='000000',font_size=16,rounded=False):
        sh=self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,*self.box(bbox))
        sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill);sh.line.color.rgb=RGBColor.from_string(line)
        self.set_text(sh,text,font_size);return self.register(sh,id,'shape')
    def text(self,id,bbox,text,size=16):
        sh=self.slide.shapes.add_textbox(*self.box(bbox));self.set_text(sh,text,size);return self.register(sh,id,'text')
    @staticmethod
    def set_text(sh,text,size):
        from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
        tf=sh.text_frame;tf.clear();tf.word_wrap=False
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        tf.vertical_anchor=MSO_ANCHOR.MIDDLE
        for i,line in enumerate(text.split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.alignment=PP_ALIGN.CENTER
            p.space_before=p.space_after=Pt(0)
            for r in p.runs:r.font.name='Arial';r.font.size=Pt(size);r.font.color.rgb=RGBColor(0,0,0)
    def edge(self,id,points,color='000000',width=1.5,start_arrow=False,end_arrow=True):
        for i,(p,q) in enumerate(zip(points,points[1:])):
            sh=self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,*self.box([p[0],p[1],q[0],q[1]]))
            sh.line.color.rgb=RGBColor.from_string(color);sh.line.width=Pt(width)
            ln=sh._element.spPr.get_or_add_ln()
            for tag,active in [('headEnd',start_arrow and i==0),('tailEnd',end_arrow and i==len(points)-2)]:
                e=OxmlElement('a:'+tag);e.set('type','triangle' if active else 'none');ln.append(e)
            self.register(sh,f'{id}:segment:{i}','connector')
    def save(self,pptx,map_path):
        self.prs.save(pptx)
        with open(map_path,'w') as f:json.dump(self.objects,f,ensure_ascii=False,indent=2)
